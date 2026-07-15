from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Degenerate-detection floors for a text line's own tight width/height,
# before inset is added -- e.g. a single punctuation mark or a detection
# jitter artifact -- not a meaningful visual minimum otherwise. Scaled by the
# same 4/3 ratio as PipelineConfig.slide_w_in/slide_h_in's own default
# (originally 0.2in/0.15in, tuned against the older 10x5.625in slide size).
MIN_TIGHT_W_IN = 0.2 * 4 / 3
MIN_TIGHT_H_IN = 0.15 * 4 / 3


def add_text_box(slide, t, W, H, config):
    rotation_deg = t.get("rotation_deg", 0.0)
    render_geom = t.get("render_geom")

    # (top, right, bottom, left), same axis order as config.text_inset_pt --
    # see its docstring for why the box is grown by this rather than kept
    # margin-free with slack baked into its raw width/height.
    inset_top_in, inset_right_in, inset_bottom_in, inset_left_in = (
        v / 72 for v in config.text_inset_pt
    )

    if rotation_deg and render_geom:
        # render_geom is the text's own tilted quad (center + true edge-length
        # width/height, in px) rather than px_bbox's axis-aligned box, which is
        # always larger than -- and wouldn't visually match -- rotated text.
        # PowerPoint rotates a shape around its own center, so the box is
        # placed/sized here as if unrotated (centered on the quad's center)
        # and rotation is applied last.
        cx, cy, w_px, h_px = render_geom
        w_in = max(MIN_TIGHT_W_IN, w_px / W * config.slide_w_in) + inset_left_in + inset_right_in
        h_in = max(MIN_TIGHT_H_IN, h_px / H * config.slide_h_in) + inset_top_in + inset_bottom_in
        x_in = cx / W * config.slide_w_in - w_in / 2
        y_in = cy / H * config.slide_h_in - h_in / 2
    else:
        x0, y0, x1, y1 = t["px_bbox"]
        tight_w_in = (x1 - x0) / W * config.slide_w_in
        tight_h_in = (y1 - y0) / H * config.slide_h_in
        x_in = x0 / W * config.slide_w_in - inset_left_in
        y_in = y0 / H * config.slide_h_in - inset_top_in
        w_in = max(MIN_TIGHT_W_IN, tight_w_in) + inset_left_in + inset_right_in
        h_in = max(MIN_TIGHT_H_IN, tight_h_in) + inset_top_in + inset_bottom_in

    box = slide.shapes.add_textbox(Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in))
    if rotation_deg:
        box.rotation = rotation_deg
    tf = box.text_frame
    tf.word_wrap = True
    # Vertically centering (rather than the default top anchor) matters
    # because box height above is deliberately tight around the *detected
    # glyph ink* (px_bbox height + inset), not around a full text-line box --
    # a font's default "single" line spacing includes its own ascent/descent
    # metrics on top of that ink height (often 30%+ more, per the font's own
    # OS/2 table, regardless of how tall the glyphs actually render), so top-
    # anchored text overflows downward past the box's bottom edge by that
    # difference. On small detections -- e.g. the label inside a compact
    # pill-shaped button -- that overflow is large relative to the box and
    # visibly pushes the label below the pill entirely. Centering distributes
    # that same overflow evenly above/below instead of dumping it all below,
    # keeping the visible glyphs close to the detected line's true vertical
    # center regardless of the substituted font's exact metrics.
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    inset_top_pt, inset_right_pt, inset_bottom_pt, inset_left_pt = config.text_inset_pt
    tf.margin_top = Pt(inset_top_pt)
    tf.margin_right = Pt(inset_right_pt)
    tf.margin_bottom = Pt(inset_bottom_pt)
    tf.margin_left = Pt(inset_left_pt)
    paragraph = tf.paragraphs[0]
    run = paragraph.add_run()
    run.text = t["text"]
    run.font.size = Pt(t["font_size_pt"])
    run.font.name = t["font_name"]
    # Exact (rather than the default "single", i.e. font-metric-driven)
    # line spacing shrinks that same ascent/descent overflow at the source,
    # on top of (not instead of) centering above -- single spacing on some
    # fonts (e.g. this document family's CJK font) reserves up to ~1.3x the
    # em size even for one line, most of it never touched by this box's
    # actual single-run content.
    paragraph.line_spacing = Pt(t["font_size_pt"])
    try:
        run.font.color.rgb = RGBColor.from_string(t["color"])
    except Exception:
        run.font.color.rgb = RGBColor(0, 0, 0)
